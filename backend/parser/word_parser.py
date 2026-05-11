"""Word document parser for pre-assistant."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Literal

logger = logging.getLogger(__name__)


try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None  # type: ignore


from backend.parser.paper_model import ParsedPaper, PaperSection


class WordParser:
    """Parse Word (.docx) documents into a ParsedPaper structure."""

    def __init__(self) -> None:
        if DocxDocument is None:
            raise ImportError(
                "python-docx is not installed. Install it with: pip install python-docx"
            )

    async def parse(self, file_path: Path, _project_dir: Path) -> ParsedPaper:
        """Parse a Word document and return a ParsedPaper."""
        logger.info(f"Parsing Word document: {file_path}")
        doc = DocxDocument(file_path)

        title = self._extract_title(doc)
        authors: list[str] = []
        abstract = ""
        sections: list[PaperSection] = []
        references: list[str] = []

        current_section: PaperSection | None = None
        current_content_lines: list[str] = []

        def flush_section() -> None:
            nonlocal current_section, current_content_lines
            if current_section and current_content_lines:
                current_section.content = "\n".join(current_content_lines).strip()
            current_content_lines = []

        def heading_level(style_name: str | None) -> int:
            if not style_name:
                return 0
            s = style_name.lower()
            if s in ("title", "heading 1", "heading1"):
                return 1
            if s in ("heading 1", "heading1"):
                return 1
            if s in ("heading 2", "heading2"):
                return 2
            if s in ("heading 3", "heading3"):
                return 3
            return 0

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            level = heading_level(para.style.name if para.style else None)

            if level == 1 and not title:
                title = text
                flush_section()
                if current_section:
                    sections.append(current_section)
                current_section = PaperSection(title=text, level=1)
                current_content_lines = []
            elif level >= 2:
                flush_section()
                if current_section:
                    sections.append(current_section)
                current_section = PaperSection(title=text, level=level)
                current_content_lines = []
            else:
                # Body text — detect abstract and references
                text_lower = text.lower()
                if not abstract and ("abstract" in text_lower or "摘要" in text):
                    abstract = text
                elif "reference" in text_lower or "参考文献" in text:
                    flush_section()
                    if current_section:
                        sections.append(current_section)
                    current_section = PaperSection(title="References", level=1)
                    current_content_lines = []
                elif current_section and current_section.title.lower() in ("references", "参考文献"):
                    references.append(text)
                elif current_section:
                    current_content_lines.append(text)
                elif sections and sections[-1].title == "References":
                    references.append(text)
                else:
                    # No section yet — create a default body section
                    if not sections and not current_section:
                        current_section = PaperSection(title="Content", level=1)
                    if current_section:
                        current_content_lines.append(text)

        flush_section()
        if current_section:
            sections.append(current_section)

        paper = ParsedPaper(
            title=title or "Untitled Document",
            authors=authors,
            abstract=abstract,
            sections=sections,
            references=references,
            source_type="docx",
        )
        logger.info(
            f"Word parsed: title={paper.title}, sections={len(paper.sections)}, "
            f"abstract={len(paper.abstract)} chars, references={len(paper.references)}"
        )
        return paper

    def _extract_title(self, doc: "DocxDocument") -> str:
        """Extract title from the document's title paragraph or first Heading 1."""
        if doc.core_properties.title:
            return doc.core_properties.title

        for para in doc.paragraphs:
            if para.style and para.style.name in ("Title", "Heading 1", "Heading1"):
                text = para.text.strip()
                if text:
                    return text

        # Fallback: first non-empty paragraph
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                return text
        return "Untitled Document"